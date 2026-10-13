from __future__ import annotations

import io
import logging
import re
import threading
import time
from dataclasses import dataclass, field
from datetime import date
from typing import Iterable

import requests
from openpyxl import load_workbook

from config import settings

log = logging.getLogger(__name__)

DAY_TO_SHEET = {
    "понедельник": "ПН",
    "пн": "ПН",
    "вторник": "ВТ",
    "вт": "ВТ",
    "среда": "СР",
    "ср": "СР",
    "четверг": "ЧТ",
    "чт": "ЧТ",
    "пятница": "ПТ",
    "пт": "ПТ",
    "суббота": "СБ",
    "сб": "СБ",
    "воскресенье": "ВС",
    "вс": "ВС",
}

REPLACEMENT_SHEETS = {
    "ПН": "Замены Пн",
    "ВТ": "Замены Вт",
    "СР": "Замены Ср",
    "ЧТ": "Замены Чт",
    "ПТ": "Замены Пт",
    "СБ": "Замены Сб",
}

RUS_WEEKDAYS = {
    "ПН": "Понедельник",
    "ВТ": "Вторник",
    "СР": "Среда",
    "ЧТ": "Четверг",
    "ПТ": "Пятница",
    "СБ": "Суббота",
    "ВС": "Воскресенье",
}

DISTANCE_SHEET = "Дистанционное обучение"

DAY_COLS = ("ПН", "ВТ", "СР", "ЧТ", "ПТ", "СБ")


@dataclass
class Lesson:
    number: int
    teacher: str
    cab: str
    day_sheet: str
    week_note: str = ""


@dataclass
class Replacement:
    day_sheet: str
    day_title: str
    group: str
    para: str
    removed: dict
    assigned: dict


@dataclass
class Distance:
    group: str
    link: str
    curator: str


@dataclass
class Timetable:
    loaded_at: float = 0.0
    lessons_by_group: dict = field(default_factory=dict)
    replacements: list = field(default_factory=list)
    distance: dict = field(default_factory=dict)
    date_by_day: dict = field(default_factory=dict)
    all_groups: set = field(default_factory=set)


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s.strip().upper())


def _tokens(s: str) -> list[str]:
    return [_norm(x) for x in s.split() if x]


def _num(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _date_from_header(header: str) -> str:
    m = re.search(r"(\d{2}\.\d{2}\.\d{4})", header or "")
    up = _norm(header)
    week = ""
    if "ВЕРХНЯЯ" in up:
        week = "верхняя"
    elif "НИЖНЯЯ" in up:
        week = "нижняя"
    if not m:
        return week
    return f"{m.group(1)} ({week})".strip()


class _Cache:
    def __init__(self) -> None:
        self._lock = threading.Lock()
        self._data: Timetable | None = None

    def get(self) -> Timetable:
        with self._lock:
            stale = (
                self._data is None
                or (time.time() - self._data.loaded_at) > settings.CACHE_TTL_SECONDS
            )
            if stale:
                try:
                    self._data = self._fetch()
                except Exception:
                    log.exception("Не удалось обновить расписание")
                    if self._data is None:
                        raise
            return self._data

    def _fetch(self) -> Timetable:
        log.info("Скачиваю таблицу: %s", settings.spreadsheet_url)
        resp = requests.get(settings.spreadsheet_url, timeout=settings.REQUEST_TIMEOUT)
        resp.raise_for_status()
        bio = io.BytesIO(resp.content)
        wb = load_workbook(bio, read_only=True, data_only=True)
        tt = Timetable(loaded_at=time.time())
        try:
            _parse_all(wb, tt)
        finally:
            wb.close()
        return tt


CACHE = _Cache()


def get_timetable() -> Timetable:
    return CACHE.get()


def _parse_all(wb, tt: Timetable) -> None:
    for sheet_name in DAY_COLS:
        if sheet_name not in wb.sheetnames:
            continue
        _parse_schedule_sheet(wb[sheet_name], sheet_name, tt)

    for day_sheet, repl_sheet in REPLACEMENT_SHEETS.items():
        if repl_sheet in wb.sheetnames:
            tt.replacements.extend(_parse_replacement_sheet(wb[repl_sheet], day_sheet))

    if DISTANCE_SHEET in wb.sheetnames:
        tt.distance = _parse_distance_sheet(wb[DISTANCE_SHEET])

    for g, _ in tt.lessons_by_group:
        tt.all_groups.add(g)
    for g in tt.distance:
        tt.all_groups.add(g)
    for r in tt.replacements:
        tt.all_groups.update(_tokens(r.group))


def _parse_schedule_sheet(ws, sheet_name: str, tt: Timetable) -> None:
    rows = list(ws.iter_rows())
    if not rows:
        return
    tt.date_by_day[sheet_name] = _date_from_header(_num(rows[0][0].value))

    header_idx = -1
    header_row = None
    for i, row in enumerate(rows):
        vals = [_norm(_num(c.value)) for c in row]
        if "ГРУППА" in vals:
            header_idx = i
            header_row = row
            break
    if header_idx < 0:
        return

    num_row = rows[header_idx - 1] if header_idx - 1 >= 0 else None
    pair_cols = []
    j = 0
    while j < len(header_row):
        if _norm(_num(header_row[j].value)) == "ГРУППА":
            lesson_no = None
            if num_row is not None and j < len(num_row):
                raw = _num(num_row[j].value)
                lesson_no = int(raw) if raw.isdigit() else None
            pair_cols.append((lesson_no, j, j + 1))
            j += 2
        else:
            j += 1

    for row in rows[header_idx + 1:]:
        teacher = _num(row[0].value)
        if not teacher or _norm(teacher) in ("ГРУППА", "КУРАТОР", "ПРЕПОДАВАТЕЛЬ"):
            continue
        for lesson_no, g_idx, c_idx in pair_cols:
            if g_idx >= len(row):
                continue
            grp_val = _num(row[g_idx].value)
            if not grp_val:
                continue
            cab = _num(row[c_idx].value) if c_idx < len(row) else ""
            for token in _tokens(grp_val):
                key = (token, sheet_name)
                tt.lessons_by_group.setdefault(key, []).append(
                    Lesson(
                        number=lesson_no or 0,
                        teacher=teacher,
                        cab=cab,
                        day_sheet=sheet_name,
                        week_note=tt.date_by_day.get(sheet_name, ""),
                    )
                )


def _parse_replacement_sheet(ws, day_sheet: str) -> list[Replacement]:
    rows = list(ws.iter_rows())
    out: list[Replacement] = []
    day_title = RUS_WEEKDAYS.get(day_sheet, day_sheet)

    start_idx = -1
    for i, row in enumerate(rows):
        vals = [_norm(_num(c.value)) for c in row]
        if vals and vals[0] == "ГРУППА":
            start_idx = i
            break
    if start_idx < 0:
        return out

    current_group = ""
    for row in rows[start_idx + 1:]:
        vals = [_num(c.value) for c in row]
        grp = vals[0] if vals and vals[0] else ""
        if grp and _norm(grp) not in ("ДИСЦИПЛИНА", "ПРЕПОДАВАТЕЛЬ", "АУД."):
            current_group = grp
        if not current_group:
            continue
        para = vals[1] if len(vals) > 1 else ""
        removed = {
            "disc": vals[2] if len(vals) > 2 else "",
            "teacher": vals[3] if len(vals) > 3 else "",
            "cab": vals[4] if len(vals) > 4 else "",
        }
        assigned = {
            "disc": vals[5] if len(vals) > 5 else "",
            "teacher": vals[6] if len(vals) > 6 else "",
            "cab": vals[7] if len(vals) > 7 else "",
        }
        if not para and not removed["disc"] and not assigned["disc"]:
            continue
        out.append(
            Replacement(
                day_sheet=day_sheet,
                day_title=day_title,
                group=current_group,
                para=para,
                removed=removed,
                assigned=assigned,
            )
        )
    return out


def _parse_distance_sheet(ws) -> dict[str, Distance]:
    out = {}
    started = False
    for row in ws.iter_rows():
        vals = [c.value for c in row]
        if not vals or vals[0] is None:
            continue
        norm0 = _norm(_num(vals[0]))
        if norm0 == "НАИМЕНОВАНИЕ ГРУППЫ" or _norm(vals[1] or "") == "ССЫЛКА НА ВСТРЕЧУ С ГРУППОЙ":
            started = True
            continue
        if not started:
            continue
        group = _num(vals[0])
        link = _num(vals[1]) if len(vals) > 1 else ""
        curator = _num(vals[2]) if len(vals) > 2 else ""
        if group and link:
            for t in _tokens(group):
                out[t] = Distance(group=t, link=link, curator=curator)
    return out


def find_group_lessons(tt: Timetable, group: str, day_sheets: Iterable[str]) -> list[Lesson]:
    target = _norm(group)
    result = []
    for sheet in day_sheets:
        result.extend(tt.lessons_by_group.get((target, sheet), []))
    seen = set()
    uniq = []
    for l in sorted(result, key=lambda x: (x.day_sheet, x.number)):
        k = (l.day_sheet, l.number, l.teacher, l.cab)
        if k not in seen:
            seen.add(k)
            uniq.append(l)
    return uniq


def guess_groups(tt: Timetable, partial: str) -> list[str]:
    p = _norm(partial)
    if not p:
        return []
    return [g for g in sorted(tt.all_groups) if p in _norm(g)]


def get_distance(tt: Timetable, group: str) -> Distance | None:
    return tt.distance.get(_norm(group))


def today_sheet() -> str:
    wd = date.today().isoweekday()
    return {1: "ПН", 2: "ВТ", 3: "СР", 4: "ЧТ", 5: "ПТ", 6: "СБ", 7: "ВС"}[wd]


def normalize_day(raw: str) -> str | None:
    n = _norm(raw)
    if n.lower() in DAY_TO_SHEET:
        return DAY_TO_SHEET[n.lower()]
    stripped = n.rstrip(".").lower()
    if stripped in DAY_TO_SHEET:
        return DAY_TO_SHEET[stripped]
    if n.isdigit() and 1 <= int(n) <= 6:
        return DAY_COLS[int(n) - 1]
    return None